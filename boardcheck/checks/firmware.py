"""Firmware against the schematic.

ADC channel maps. Firmware often names each telemetry channel in a C enum
whose position selects the ADC and its input (for CM-03545's tlm_adc.c:
input = index & 7, chip select = index >> 3). The check follows each entry
from the source files to the board: the SPI select the firmware asserts is
a SmartDesign pin (SPISS[n] of the SPI core), connected to a top-level
port, constrained to an FPGA ball, wired (through series resistors) to one
ADC's chip select; the channel number picks that ADC's input pin, whose net
names the signal the board actually feeds it. That name must match the
firmware's. Configure under `firmware.adc_channel_maps` (paths relative to
the config file, as for `fpga`).
"""

import os
import re

from . import ERROR, INFO, WARNING, Finding, check
from .levels import signals
from .pins import _pin_types
from ..model import natural_key


def parse_enum(path, enum_name):
    """[(name, value)] of a C typedef enum, honouring explicit values."""
    with open(path, encoding="utf-8", errors="replace") as f:
        text = f.read()
    text = re.sub(r"/\*.*?\*/", " ", text, flags=re.S)
    text = re.sub(r"//[^\n]*", " ", text)
    m = re.search(r"typedef\s+enum\s*\w*\s*\{([^}]*)\}\s*" + re.escape(enum_name) + r"\s*;", text)
    if not m:
        m = re.search(r"enum\s+" + re.escape(enum_name) + r"\s*\{([^}]*)\}", text)
    if not m:
        raise ValueError(f"enum {enum_name} not found in {path}")
    out, value = [], 0
    for item in (x.strip() for x in m.group(1).split(",")):
        if not item:
            continue
        name, _, expr = item.partition("=")
        if expr.strip():
            value = int(expr.strip(), 0)
        out.append((name.strip(), value))
        value += 1
    return out


class ChannelMap:
    def __init__(self, ctx, spec):
        self.ctx = ctx
        self.spec = spec
        self.name = spec.get("name", spec.get("enum", "adc map"))
        self.problems = []
        self.entries = []           # (fw name, stem, select, channel, adc pin or None, net or None)
        self.adcs = {}              # select -> ADC component
        self._build()

    def _path(self, p):
        return os.path.normpath(os.path.join(self.ctx.config.base_dir, os.path.expanduser(p)))

    def _build(self):
        spec, ctx = self.spec, self.ctx
        path = self._path(spec["enum_file"])
        if not os.path.isfile(path):
            self.problems.append(f"enum file not found: {path}")
            return
        try:
            enum = parse_enum(path, spec["enum"])
        except ValueError as e:
            self.problems.append(str(e))
            return
        skip = set(spec.get("skip") or [])
        per = int(spec.get("channels_per_select", 8))
        comp = ctx.design.components.get(spec["fpga"])
        f = ctx.fpga_for(comp) if comp is not None else None
        if f is None:
            self.problems.append(f"FPGA {spec['fpga']} constraints not available")
            return
        sel_re = spec["select_link"]
        by_port = {c.port: c for c in f.io.pins.values()}
        pins = {str(p.designator): p for p in comp.pins}
        by_net = {n: s for s in signals(ctx) for n in s.nets}
        cs_name = spec.get("adc_select_pin", "CS")
        in_fmt = spec.get("adc_input_pin", "IN{ch}")
        for select in sorted({v // per for n, v in enum if n not in skip}):
            link = re.compile(sel_re.format(n=select))
            ports = [p for p, links in f.io.port_links.items() if any(link.search(x) for x in links)]
            if not ports:
                self.problems.append(f"select {select}: no top-level port connects to /{link.pattern}/")
                continue
            c = by_port.get(ports[0])
            pin = pins.get(c.ball) if c else None
            sig = by_net.get(pin.net) if pin else None
            adcs = [q for q in (sig.pins if sig else []) if q.component is not comp and q.name == cs_name]
            if len(adcs) != 1:
                self.problems.append(f"select {select} ('{ports[0]}'): reaches {len(adcs)} '{cs_name}' pins "
                                     f"({', '.join(q.ref for q in adcs) or 'none'})")
                continue
            self.adcs[select] = adcs[0].component
        prefix = spec.get("strip_prefix", "")
        for name, value in enum:
            if name in skip:
                continue
            select, ch = value // per, value % per
            stem = name[len(prefix):] if prefix and name.startswith(prefix) else name
            adc = self.adcs.get(select)
            pin = None
            if adc is not None:
                pin = next((p for p in adc.pins if p.name == in_fmt.format(ch=ch)), None)
            self.entries.append((name, stem, select, ch, pin, pin.net if pin else None))

    def net_stem(self, net):
        s = net
        for rx in self.spec.get("net_strip") or []:
            s = re.sub(rx, "", s)
        return s


def _maps(ctx):
    if not hasattr(ctx, "_fw_maps"):
        ctx._fw_maps = [ChannelMap(ctx, spec) for spec in (ctx.config["firmware"]["adc_channel_maps"] or [])]
    return ctx._fw_maps


@check("FW001", "Firmware ADC channel reads a different signal than it names", ERROR)
def channel_names(ctx):
    for m in _maps(ctx):
        for p in m.problems:
            yield Finding("FW001", f"{m.name}: {p}", severity=WARNING)
        for name, stem, select, ch, pin, net in m.entries:
            if pin is None:
                continue
            if m.net_stem(net).upper() != stem.upper():
                yield Finding("FW001", f"{m.name}: firmware '{name}' reads {pin.ref} {pin.name} (select {select}), "
                                       f"which the schematic wires to '{net}'", refs=[pin.component.designator],
                              nets=[net])


@check("FW002", "ADC channel wired but unread, or read but unwired", WARNING)
def channel_coverage(ctx):
    for m in _maps(ctx):
        read = {(pin.component.designator, pin.designator) for *_, pin, _ in m.entries if pin is not None}
        for name, stem, select, ch, pin, net in m.entries:
            if m.adcs.get(select) is not None and (pin is None or len(ctx.design.nets[net].pins) < 2):
                yield Finding("FW002", f"{m.name}: firmware '{name}' reads select {select} input {ch}, which is "
                                       "not connected on the board",
                              refs=[m.adcs[select].designator])
        in_re = re.compile("^" + re.escape(m.spec.get("adc_input_pin", "IN{ch}")).replace(r"\{ch\}", r"\d+") + "$")
        for select, adc in sorted(m.adcs.items()):
            for p in sorted(adc.pins, key=lambda p: natural_key(p.name)):
                if in_re.match(p.name) and (adc.designator, p.designator) not in read \
                        and len(ctx.design.nets[p.net].pins) > 1 and not ctx.config.is_ground(p.net):
                    yield Finding("FW002", f"{m.name}: {p.ref} {p.name} is wired to '{p.net}' but no firmware "
                                           "channel reads it", refs=[adc.designator], nets=[p.net])


@check("FW003", "Firmware ADC channel map", INFO)
def channel_map(ctx):
    for m in _maps(ctx):
        rows = [f"{name}={pin.ref if pin else '?'}" for name, stem, select, ch, pin, net in m.entries]
        if rows:
            yield Finding("FW003", f"{m.name}: {len(rows)} channels traced firmware -> board "
                                   f"({', '.join(sorted({a.designator for a in m.adcs.values()}, key=natural_key))})")


# -- scaling ------------------------------------------------------------------------

def channel_scaling(ctx, m, pin):
    """(ratio, upstream net, resistor designators) for an ADC input fed
    through a series resistor from an upstream net, with an optional
    resistor to ground: ratio = Rbottom / (Rtop + Rbottom), 1.0 without a
    bottom resistor. None when the input is not wired that way or a value is
    unknown."""
    from .levels import _ohms
    cfg = ctx.config
    net = ctx.design.nets[pin.net]
    top, bottom = [], []
    for q in net.pins:
        comp = q.component
        if ctx.kind(comp) != "resistor" or len(comp.pins) != 2:
            continue
        other = next(x.net for x in comp.pins if x is not q)
        (bottom if cfg.is_ground(other) else top).append((comp, other))
    if len(top) != 1 or len(bottom) > 1:
        return None
    rt = _ohms(ctx, top[0][0])
    rb = _ohms(ctx, bottom[0][0]) if bottom else None
    if rt is None or (bottom and rb is None):
        return None
    ratio = rb / (rt + rb) if bottom else 1.0
    return ratio, top[0][1], [top[0][0].designator] + ([bottom[0][0].designator] if bottom else [])


_CURRENT = r"(ISENSE|_I$|CURRENT)"


def _stem_volts(ctx, stem):
    """Nominal volts of the rail a channel is named after ("28V0_EPS"), by
    the rail naming convention; None for current and other channels."""
    m = ctx.config._rail_re.match(stem)
    if not m or re.search(_CURRENT, stem, re.I):
        return None
    return float(f"{m.group('int')}.{m.group('frac') or '0'}")


def _reference_volts(ctx, m, adc):
    ref = m.spec.get("adc_reference_pin", "VA")
    nets = {p.net for p in adc.pins if p.name == ref}
    volts = {ctx.config.net_voltage(n) for n in nets} - {None}
    return volts.pop() if len(volts) == 1 else None


@check("FW004", "ADC channel's nominal input exceeds the ADC reference", ERROR)
def channel_full_scale(ctx):
    """A voltage channel named after a rail ("TLM_3V3_MISC") reads that
    rail through its divider; at the rail's nominal voltage the ADC input
    must stay below the ADC reference (`adc_reference_pin`)."""
    for m in _maps(ctx):
        for name, stem, select, ch, pin, net in m.entries:
            if pin is None:
                continue
            volts = _stem_volts(ctx, stem)
            sc = channel_scaling(ctx, m, pin)
            ref = _reference_volts(ctx, m, pin.component)
            if volts is None or sc is None or ref is None:
                continue
            at_pin = volts * sc[0]
            if at_pin > ref:
                yield Finding("FW004", f"{m.name}: '{name}' ({stem}, {volts:g} V nominal) reaches {pin.ref} "
                                       f"{pin.name} at {at_pin:.2f} V through {'/'.join(sc[2])} (ratio {sc[0]:.4g}), "
                                       f"above the {ref:g} V reference", refs=[pin.component.designator], nets=[net])


def _read_cal(path, sheet, columns):
    """[(signal, gain, offset)] from a calibration spreadsheet (.xlsx or .csv)."""
    names = [c.lower() for c in columns]
    rows = []
    if path.lower().endswith(".csv"):
        import csv
        with open(path, newline="", encoding="utf-8") as f:
            rows = [tuple(r) for r in csv.reader(f)]
    else:
        import openpyxl     # optional dependency: only needed for .xlsx calibration files
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        rows = list(wb[sheet].iter_rows(values_only=True))
    header = [str(h or "").strip().lower() for h in rows[0]]
    idx = [header.index(n) for n in names]
    return [tuple(r[i] for i in idx) for r in rows[1:] if r and r[idx[0]] not in (None, "")]


def _cal_rows(ctx, m):
    spec = m.spec.get("calibration")
    if not spec:
        return None, None
    path = m._path(spec["file"])
    if not os.path.isfile(path):
        return None, f"calibration file not found: {path}"
    try:
        return _read_cal(path, spec.get("sheet"), spec.get("columns", ["Signal", "Gain", "Offset"])), None
    except ImportError:
        return None, "reading .xlsx calibration files needs openpyxl (pip install openpyxl)"
    except (KeyError, ValueError) as e:
        return None, f"calibration file {os.path.basename(path)}: {e}"


def _placeholders(rows):
    """True when every row is gain 1, offset 0: a table never filled in."""
    try:
        return bool(rows) and all(float(r[1]) == 1.0 and float(r[2] or 0) == 0.0 for r in rows)
    except (TypeError, ValueError):
        return False


@check("FW005", "Calibration gain differs from the board's scaling", WARNING)
def calibration_gains(ctx):
    """For each voltage channel, the gain the board implies (reference / 2^bits
    / divider ratio: millivolts at the rail per count, the scaling
    tlm_adc.c's comment describes) against the calibration file's gain, by
    position in the firmware enum. Current channels are not computed."""
    for m in _maps(ctx):
        rows, problem = _cal_rows(ctx, m)
        if problem:
            yield Finding("FW005", f"{m.name}: {problem}", severity=INFO)
        if not rows:
            continue
        if _placeholders(rows):
            yield Finding("FW005", f"{m.name}: every row of {os.path.basename(m.spec['calibration']['file'])} is "
                                   f"gain 1, offset 0: no channel is calibrated, so telemetry reads raw counts")
            continue
        bits = int(m.spec.get("adc_bits", 12))
        tol = float(m.spec.get("calibration", {}).get("tolerance", 0.05))
        off = []
        for (name, stem, select, ch, pin, net), row in zip(m.entries, rows):
            volts = _stem_volts(ctx, stem)
            sc = channel_scaling(ctx, m, pin) if pin is not None else None
            ref = _reference_volts(ctx, m, pin.component) if pin is not None else None
            if volts is None or sc is None or ref is None:
                continue
            expected = ref * 1000.0 / (1 << bits) / sc[0]
            try:
                gain = float(row[1])
            except (TypeError, ValueError):
                continue
            if abs(gain - expected) > tol * expected:
                off.append(f"{name} {gain:g} (board: {expected:.4g} mV/count, ratio {sc[0]:.4g})")
        if off:
            yield Finding("FW005", f"{m.name}: {len(off)} voltage channel gain(s) in "
                                   f"{os.path.basename(m.spec['calibration']['file'])} differ from the board's "
                                   f"scaling by more than {tol:.0%}: " + "; ".join(off[:8])
                                   + (f"; ... {len(off) - 8} more" if len(off) > 8 else ""))


@check("FW006", "Calibration file signals differ from the firmware enum", WARNING)
def calibration_names(ctx):
    """The calibration table is loaded by position (cal.c copies it into
    TLM_Cal_t[] indexed by the enum), so its rows must list the enum's
    signals in the same order. Names are not used, so stray trailing
    punctuation ('TLM_1V2_8GB,') is ignored."""
    for m in _maps(ctx):
        rows, _ = _cal_rows(ctx, m)
        if not rows:
            continue
        names = [e[0] for e in m.entries]
        cal = [str(r[0]).strip().rstrip(",;").strip() for r in rows]
        if len(cal) != len(names):
            yield Finding("FW006", f"{m.name}: calibration file has {len(cal)} rows, the enum {len(names)} signals")
        bad = [(i, c, n) for i, (c, n) in enumerate(zip(cal, names)) if c != n]
        if bad:
            yield Finding("FW006", f"{m.name}: calibration rows that do not name the enum signal at their position: "
                                   + "; ".join(f"row {i + 1} '{c}' (enum: {n})" for i, c, n in bad[:10])
                                   + (f"; ... {len(bad) - 10} more" if len(bad) > 10 else ""))


# -- PWM / DAC full scale ---------------------------------------------------------

def parse_define(path, macro):
    """The numeric value of `#define MACRO value` in a C header, or None."""
    with open(path, encoding="utf-8", errors="replace") as f:
        for line in f:
            m = re.match(r"\s*#\s*define\s+" + re.escape(macro) + r"\s+\(?\s*([-+]?(?:0x[0-9a-fA-F]+|\d+(?:\.\d*)?))", line)
            if m:
                return float(int(m.group(1), 0)) if m.group(1).lower().startswith(("0x", "-0x")) else float(m.group(1))
    return None


def pwm_full_scale(ctx, f, pin):
    """(volts at the load at 100% duty, load pins, resistors, rail) for an
    FPGA PWM output filtered by a series resistor into a node with resistors
    to ground. The high level is the pin's bank rail. None when the path is
    not that shape or a value is unknown."""
    from .levels import _ohms
    cfg = ctx.config
    bank = f.bank(pin)
    nets = f.bank_supply_nets(bank) if bank is not None else set()
    rails = {cfg.net_voltage(n) for n in nets} - {None}
    if len(rails) != 1:
        return None
    v_high = rails.pop()
    first = ctx.design.nets[pin.net]
    series = [q for q in first.pins if q.component is not pin.component and ctx.kind(q.component) == "resistor"
              and len(q.component.pins) == 2]
    if len(series) != 1:
        return None
    rs_comp = series[0].component
    node = next(p.net for p in rs_comp.pins if p is not series[0])
    if cfg.is_ground(node) or cfg.is_rail(node):
        return None
    rs = _ohms(ctx, rs_comp)
    downs = []
    for q in ctx.design.nets[node].pins:
        c = q.component
        if c is rs_comp or ctx.kind(c) != "resistor" or len(c.pins) != 2:
            continue
        other = next(p.net for p in c.pins if p is not q)
        if not cfg.is_ground(other):
            return None         # another source or series path: not a simple filter
        downs.append(c)
    ohms = [_ohms(ctx, c) for c in downs]
    if rs is None or any(o is None for o in ohms):
        return None
    if ohms:
        rb = 1.0 / sum(1.0 / o for o in ohms)
        ratio = rb / (rs + rb)
    else:
        ratio = 1.0
    loads = [q for q in ctx.design.nets[node].pins if ctx.kind(q.component) not in ("resistor", "capacitor")]
    return v_high * ratio, loads, [rs_comp.designator] + [c.designator for c in downs], v_high


@check("FW007", "Firmware full-scale constant differs from the board", ERROR)
def pwm_constants(ctx):
    """firmware.pwm_outputs: a firmware constant stating the voltage a PWM
    output produces at its load at 100% duty (e.g. PWM_VREF_mV) against the
    board: the FPGA bank rail through the RC filter's divider."""
    for spec in ctx.config["firmware"]["pwm_outputs"] or []:
        path = os.path.normpath(os.path.join(ctx.config.base_dir, os.path.expanduser(spec["constant_file"])))
        name = spec.get("name", spec["constant"])
        if not os.path.isfile(path):
            yield Finding("FW007", f"{name}: firmware file not found: {path}", severity=WARNING)
            continue
        value = parse_define(path, spec["constant"])
        if value is None:
            yield Finding("FW007", f"{name}: #define {spec['constant']} not found in {os.path.basename(path)}",
                          severity=WARNING)
            continue
        volts_fw = value * {"mV": 1e-3, "V": 1.0}[spec.get("constant_unit", "mV")]
        comp = ctx.design.components.get(spec["fpga"])
        f = ctx.fpga_for(comp) if comp is not None else None
        if f is None:
            continue
        by_port = {c.port: c for c in f.io.pins.values()}
        pins = {str(p.designator): p for p in comp.pins}
        tol = float(spec.get("tolerance", 0.05))
        for port in spec["ports"]:
            c = by_port.get(port)
            pin = pins.get(c.ball) if c else None
            fs = pwm_full_scale(ctx, f, pin) if pin else None
            if fs is None:
                yield Finding("FW007", f"{name}: could not work out the full scale of '{port}' "
                                       "(expected FPGA pin -> series resistor -> node with resistors to ground)",
                              severity=INFO, refs=[spec["fpga"]])
                continue
            volts, loads, rs, v_high = fs
            if abs(volts - volts_fw) > tol * volts_fw:
                where = ", ".join(q.ref + (f" {q.name}" if q.name else "") for q in loads) or "its filter node"
                yield Finding("FW007", f"{name}: {spec['constant']} = {value:g} {spec.get('constant_unit', 'mV')} "
                                       f"({os.path.basename(path)}) but '{port}' reaches {where} at "
                                       f"{volts:.3g} V full scale ({v_high:g} V bank rail through "
                                       f"{'/'.join(rs)}); setpoints are scaled by {volts / volts_fw:.3g}",
                              refs=sorted({spec["fpga"]} | {q.component.designator for q in loads}, key=natural_key))


@check("FW008", "Firmware limit beyond what the board can produce", WARNING, needs_partsdb=True)
def pwm_limits(ctx):
    """firmware.pwm_outputs[].limit: a firmware maximum for the quantity the
    PWM sets at its load (e.g. I_MAX_MA for a stepper driver whose current
    is VREF / KV) against the most the board can produce: the output's full
    scale divided by the load pin's characteristic (`gain`, e.g. kv),
    taken at its most favourable limit."""
    from .levels import PinLevels, _levels
    lv = _levels(ctx)
    for spec in ctx.config["firmware"]["pwm_outputs"] or []:
        lim = spec.get("limit")
        if not lim:
            continue
        path = os.path.normpath(os.path.join(ctx.config.base_dir, os.path.expanduser(spec["constant_file"])))
        if not os.path.isfile(path):
            continue
        value = parse_define(path, lim["constant"])
        if value is None:
            yield Finding("FW008", f"{spec.get('name')}: #define {lim['constant']} not found", severity=INFO)
            continue
        scale = {"mA": 1e-3, "A": 1.0, "mV": 1e-3, "V": 1.0}[lim.get("unit", "mA")]
        comp = ctx.design.components.get(spec["fpga"])
        f = ctx.fpga_for(comp) if comp is not None else None
        if f is None:
            continue
        by_port = {c.port: c for c in f.io.pins.values()}
        pins = {str(p.designator): p for p in comp.pins}
        for port in spec["ports"]:
            c = by_port.get(port)
            pin = pins.get(c.ball) if c else None
            fs = pwm_full_scale(ctx, f, pin) if pin else None
            if fs is None:
                continue
            volts, loads, rs, _ = fs
            for load in loads:
                chars = ctx.partsdb.characteristics(load.component.part_number) or {}
                pp = _pin_types(ctx).part_entry(load)
                if pp is None or not chars:
                    continue
                pl = PinLevels(load, chars, key=pp.key)
                g_min = lv.value(pl, lim["gain"], "min", "low")[0]
                g_typ = next((r.get("typ") for _, char in lv._tables(pl, lim["gain"]) for r in char.get("rows", [])
                              if "typ" in r), None)
                if not g_min:
                    continue
                best = volts / g_min
                if value * scale > best * 1.0001:
                    typ = f", {volts / g_typ:.3g} typical" if g_typ else ""
                    yield Finding("FW008", f"{spec.get('name')}: {lim['constant']} = {value:g} {lim.get('unit', 'mA')} "
                                           f"but '{port}' reaches {load.ref} {load.name} at {volts:.3g} V full scale, "
                                           f"which gives at most {best:.3g} {lim.get('result_unit', 'A')} "
                                           f"({lim['gain']} {g_min:g} minimum{typ})",
                                  refs=sorted({spec["fpga"], load.component.designator}, key=natural_key))


# -- current channels ------------------------------------------------------------------

class CurrentChannel:
    """A current telemetry channel's scaling, worked out from the board: the
    shunt, the amplifier gain (ADC volts per shunt volt) and the ADC input
    with no current flowing (the amplifier's reference offset)."""

    def __init__(self, entry, gain, shunt, ohms, zero, ref, bits):
        self.entry = entry
        self.gain = gain
        self.shunt = shunt
        self.ohms = ohms
        self.zero = zero            # volts at the ADC input at zero current, None if unknown
        self.ref = ref              # ADC reference volts, None if unknown
        self.bits = bits

    @property
    def amps_per_count(self):
        return self.ref / (1 << self.bits) / (self.gain * self.ohms) if self.ref else None

    @property
    def zero_counts(self):
        return self.zero * (1 << self.bits) / self.ref if self.ref and self.zero is not None else None

    @property
    def full_scale_amps(self):
        return (self.ref - (self.zero or 0.0)) / (self.gain * self.ohms) if self.ref else None


def current_channels(ctx, m):
    """[(entry, CurrentChannel or None)] for the map's current channels
    (names matching `current_pattern`), each traced from its ADC input back
    through resistors and op-amps to one shunt (at most `shunt_max_ohms`)."""
    from ..analog import current_gain, zero_output
    pattern = re.compile(m.spec.get("current_pattern", _CURRENT), re.I)
    shunt_max = float(m.spec.get("shunt_max_ohms", 0.1))
    bits = int(m.spec.get("adc_bits", 12))
    out = []
    for entry in m.entries:
        name, stem, select, ch, pin, net = entry
        if pin is None or not pattern.search(stem):
            continue
        r = current_gain(ctx, net, shunt_max)
        if r is None or r[0] == 0:
            out.append((entry, None))
            continue
        gain, shunt, ohms, network = r
        out.append((entry, CurrentChannel(entry, gain, shunt, ohms, zero_output(network, net),
                                          _reference_volts(ctx, m, pin.component), bits)))
    return out


@check("FW009", "Current channel scaling from the board", INFO)
def current_scaling(ctx):
    """For each current channel: the shunt, the amplifier gain, and what that
    makes one ADC count, the zero-current reading and the full-scale current.
    The calibration gain for a current channel should be the A/count (or
    mA/count) figure, in whatever unit the firmware reports."""
    for m in _maps(ctx):
        for (name, stem, select, ch, pin, net), cc in current_channels(ctx, m):
            if cc is None:
                yield Finding("FW009", f"{m.name}: '{name}' ({pin.ref} {pin.name}): could not trace '{net}' back "
                                       "through resistors and op-amps to a single shunt",
                              refs=[pin.component.designator], nets=[net])
                continue
            text = f"{m.name}: '{name}' reads shunt {cc.shunt} ({cc.ohms * 1000:g} mOhm) with gain {cc.gain:.4g}"
            if cc.ref:
                text += (f": {cc.amps_per_count * 1000:.4g} mA/count, "
                         + (f"zero current at {cc.zero:.3g} V ({cc.zero_counts:.0f} counts), " if cc.zero is not None
                            else "")
                         + f"full scale {cc.full_scale_amps:.3g} A")
            yield Finding("FW009", text, refs=sorted({cc.shunt, pin.component.designator}, key=natural_key),
                          nets=[net])


@check("FW010", "Calibration offset differs from the board's zero-current reading", WARNING)
def current_offsets(ctx):
    """tlm_adc.c scales a reading as raw * gain + offset. When the current
    amplifier's output sits at a reference voltage with no current flowing,
    that reads as zero_counts counts, so the offset must be -gain *
    zero_counts whatever unit the gain is in. Compared by position in the
    firmware enum, as the calibration table is loaded. A table of
    placeholders is reported once, by FW005."""
    for m in _maps(ctx):
        rows, _ = _cal_rows(ctx, m)
        if not rows or _placeholders(rows):
            continue
        tol = float(m.spec.get("calibration", {}).get("tolerance", 0.05))
        index = {e[0]: i for i, e in enumerate(m.entries)}
        off = []
        refs = set()
        for entry, cc in current_channels(ctx, m):
            i = index[entry[0]]
            if cc is None or cc.zero_counts is None or i >= len(rows):
                continue
            try:
                gain, offset = float(rows[i][1]), float(rows[i][2])
            except (TypeError, ValueError):
                continue
            if gain == 0:
                continue
            implied = -offset / gain + 0.0  # counts the calibration treats as zero current
            if abs(implied - cc.zero_counts) > max(tol * cc.zero_counts, 2.0):
                off.append(f"{entry[0]} gain {gain:g} offset {offset:g} (zero current at {implied:.0f} counts; "
                           f"board: {cc.zero_counts:.0f} counts, {cc.zero:.3g} V)")
                refs.add(cc.shunt)
        if off:
            yield Finding("FW010", f"{m.name}: {len(off)} current channel offset(s) in "
                                   f"{os.path.basename(m.spec['calibration']['file'])} do not remove the "
                                   "amplifier's zero-current output: " + "; ".join(off[:8])
                                   + (f"; ... {len(off) - 8} more" if len(off) > 8 else ""),
                          refs=sorted(refs, key=natural_key))


# -- scale table ---------------------------------------------------------------------

_SCALE_ROW = (r"^\s*//\s*(?P<name>\S+)\s+(?P<kind>VOLTAGE|CURRENT)\s+IN(?P<ch>\d+)\s+CS_(?P<cs>\d+)"
              r"\s+(?P<current>\S+)\s+(?P<voltage>\S+)")


def scale_table(ctx, m):
    """{signal: (kind, select, channel, scale per count, "file:line")} from
    the map's `scale_table` (a table of per-channel scale factors, e.g.
    Camera/include/telemetry.h: `//3V3_ETH1 VOLTAGE IN5 CS_4 N/A 0.00132967`),
    or None when none is configured."""
    spec = m.spec.get("scale_table")
    if not spec:
        return None
    path = m._path(spec["file"])
    if not os.path.isfile(path):
        m.problems.append(f"scale table not found: {path}")
        return None
    rx = re.compile(spec.get("pattern", _SCALE_ROW))
    base = int(spec.get("select_base", 1))
    rows = {}
    with open(path, encoding="utf-8", errors="replace") as f:
        for i, line in enumerate(f, 1):
            r = rx.match(line)
            if not r:
                continue
            kind = r.group("kind").lower()
            try:
                scale = float(r.group(kind))
            except ValueError:
                continue
            rows[r.group("name")] = (kind, int(r.group("cs")) - base, int(r.group("ch")), scale,
                                     f"{os.path.basename(path)}:{i}")
    return rows


@check("FW014", "Scale table differs from the board", WARNING)
def scale_table_vs_board(ctx):
    """A table of per-channel scale factors (`scale_table`) against the
    board: each row's ADC input and chip select against the firmware enum's
    position for that signal, a voltage row's volts per count against the
    reference / 2^bits / divider ratio, a current row's amps per count
    against the shunt and amplifier (FW009). `tolerance` (default 2%)."""
    for m in _maps(ctx):
        rows = scale_table(ctx, m)
        if not rows:
            continue
        tol = float(m.spec["scale_table"].get("tolerance", 0.02))
        bits = int(m.spec.get("adc_bits", 12))
        currents = {e[1]: cc for e, cc in current_channels(ctx, m)}
        for name, stem, select, ch, pin, net in m.entries:
            row = rows.get(stem)
            if row is None:
                yield Finding("FW014", f"{m.name}: '{name}' has no row in the scale table")
                continue
            kind, r_sel, r_ch, scale, where = row
            if (r_sel, r_ch) != (select, ch):
                yield Finding("FW014", f"{m.name}: {where} puts '{stem}' on select {r_sel} input {r_ch}, "
                                       f"the firmware enum on select {select} input {ch}")
            if pin is None:
                continue
            if kind == "voltage":
                sc = channel_scaling(ctx, m, pin)
                ref = _reference_volts(ctx, m, pin.component)
                if sc is None or ref is None:
                    continue
                board = ref / (1 << bits) / sc[0]
                how = f"{'/'.join(sc[2])}, ratio {sc[0]:.4g}"
                unit, mult = "mV", 1e3
            else:
                cc = currents.get(stem)
                if cc is None or cc.amps_per_count is None:
                    continue
                board = cc.amps_per_count
                how = f"shunt {cc.shunt} {cc.ohms * 1e3:g} mOhm, gain {cc.gain:.4g}"
                unit, mult = "mA", 1e3
            if abs(scale - board) > tol * board:
                refs = [pin.component.designator] + (sc[2] if kind == "voltage" else [cc.shunt])
                yield Finding("FW014", f"{m.name}: {where} scales '{stem}' at {scale * mult:.5g} {unit}/count, "
                                       f"the board at {board * mult:.5g} {unit}/count ({how}): "
                                       f"{scale / board:.3g} times the board", refs=refs, nets=[net])


# -- GPIO bit maps ---------------------------------------------------------------------

def parse_gpio_header(path):
    """{group comment: [(define, bit)]} from a header of `#define NAME GPIO_<n>`
    lines, grouped under the comment line that precedes them."""
    groups, group = {}, None
    with open(path, encoding="utf-8", errors="replace") as f:
        for line in f:
            m = re.match(r"\s*//\s*(.+?)\s*$", line)
            if m:
                group = m.group(1)
                continue
            m = re.match(r"\s*#\s*define\s+(\w+)\s+GPIO_(\d+)\b", line)
            if m:
                groups.setdefault(group, []).append((m.group(1), int(m.group(2))))
    return groups


def _norm_name(s):
    return re.sub(r"[^A-Z0-9]", "", str(s).upper())


class GpioMap:
    """A firmware GPIO header traced through the FPGA's SmartDesign
    hierarchy to top-level ports, balls and nets."""

    def __init__(self, ctx, spec):
        from ..smartdesign import Hierarchy, SmartDesign
        self.ctx = ctx
        self.spec = spec
        self.name = spec.get("name", "gpio map")
        self.problems = []
        self.entries = []       # (define, group, instance path, pin, bit, trace, board pin or None)
        self.unmapped = []      # header groups no block claims
        base = ctx.config.base_dir
        path = os.path.normpath(os.path.join(base, os.path.expanduser(spec["header"])))
        if not os.path.isfile(path):
            self.problems.append(f"header not found: {path}")
            return
        groups = parse_gpio_header(path)
        fspec = (ctx.config["fpga"] or {}).get(spec["fpga"]) or {}
        tops = [p for p in fspec.get("top_level") or [] if p.lower().endswith(".tcl")]
        if not tops:
            self.problems.append(f"FPGA {spec['fpga']} has no SmartDesign top_level in the config")
            return
        top = os.path.normpath(os.path.join(base, tops[0]))
        root = spec.get("design_root")
        root = os.path.normpath(os.path.join(base, root)) if root else os.path.dirname(os.path.dirname(os.path.dirname(top)))
        if not os.path.isfile(top):
            self.problems.append(f"top level not found: {top}")
            return
        hier = Hierarchy(root, SmartDesign.read(top).name)
        comp = ctx.design.components.get(spec["fpga"])
        f = ctx.fpga_for(comp) if comp is not None else None
        by_port = {c.port: c for c in f.io.pins.values()} if f else {}
        balls = {str(p.designator): p for p in comp.pins} if comp is not None else {}
        claimed = set()
        for block in spec.get("blocks") or []:
            inst = block["instance"].split("/")
            for g in block.get("groups") or []:
                matches = [k for k in groups if k and k.startswith(g)]
                if not matches:
                    self.problems.append(f"no group '{g}' in {os.path.basename(path)}")
                for k in matches:
                    claimed.add(k)
                    for define, bit in groups[k]:
                        tr = hier.trace_up(inst, block.get("pin", "GPIO_OUT"), bit)
                        pin = None
                        if tr[0] == "port":
                            port = tr[1] if tr[2] is None else f"{tr[1]}[{tr[2]}]"
                            c = by_port.get(port)
                            pin = balls.get(c.ball) if c else None
                        self.entries.append((define, k, block["instance"], block.get("pin", "GPIO_OUT"), bit, tr, pin))
        self.unmapped = [k for k in groups if k not in claimed]
        self.used = set()
        src = spec.get("source_dir")
        if src:
            src = os.path.normpath(os.path.join(base, src))
            names = {e[0] for e in self.entries}
            for d, _, files in os.walk(src):
                for n in files:
                    p = os.path.join(d, n)
                    if n.endswith((".c", ".h")) and os.path.normpath(p) != path:
                        with open(p, encoding="utf-8", errors="replace") as fh:
                            text = fh.read()
                        self.used |= {x for x in names if re.search(r"\b" + x + r"\b", text)}

    @staticmethod
    def port_name(tr):
        return tr[1] if tr[2] is None else f"{tr[1]}[{tr[2]}]"


def _gpio_maps(ctx):
    if not hasattr(ctx, "_gpio_maps"):
        ctx._gpio_maps = [GpioMap(ctx, s) for s in (ctx.config["firmware"]["gpio_maps"] or [])]
    return ctx._gpio_maps


@check("FW011", "Firmware GPIO bit reaches an FPGA port of another name", ERROR)
def gpio_names(ctx):
    """Each `#define NAME GPIO_<n>` in the header names what bit n of its
    CoreGPIO carries. Followed up the SmartDesign hierarchy (through buffer
    and inverter macros) to the top-level port, the port's name must match
    the define's, ignoring case and punctuation."""
    for m in _gpio_maps(ctx):
        for p in m.problems:
            yield Finding("FW011", f"{m.name}: {p}", severity=WARNING)
        for define, group, inst, pin, bit, tr, bpin in m.entries:
            if tr[0] == "port" and _norm_name(define) != _norm_name(m.port_name(tr)):
                where = f", ball {bpin.designator} '{bpin.net}'" if bpin else ""
                inv = " (inverted)" if tr[3] else ""
                yield Finding("FW011", f"{m.name}: firmware '{define}' is {inst}:{pin}[{bit}], which reaches top-level "
                                       f"port '{m.port_name(tr)}'{inv}{where}", refs=[m.spec["fpga"]],
                              nets=[bpin.net] if bpin else [])


@check("FW012", "Firmware uses a GPIO bit the FPGA does not connect", WARNING)
def gpio_unconnected(ctx):
    """A define the firmware sources use (`source_dir`) whose bit the FPGA
    design ties to a constant or leaves unconnected: writes go nowhere,
    reads return a constant."""
    for m in _gpio_maps(ctx):
        for define, group, inst, pin, bit, tr, bpin in m.entries:
            if define not in m.used:
                continue
            if tr[0] == "constant":
                yield Finding("FW012", f"{m.name}: firmware uses '{define}' ({inst}:{pin}[{bit}]), which the FPGA "
                                       f"design ties to {tr[1]}", refs=[m.spec["fpga"]])
            elif tr[0] == "open":
                yield Finding("FW012", f"{m.name}: firmware uses '{define}' ({inst}:{pin}[{bit}]), which connects to "
                                       "nothing in the FPGA design", refs=[m.spec["fpga"]])


@check("FW013", "Firmware GPIO map", INFO)
def gpio_map(ctx):
    for m in _gpio_maps(ctx):
        if not m.entries:
            continue
        board = [e for e in m.entries if e[6] is not None]
        internal = [e[0] for e in m.entries if e[5][0] == "internal"]
        idle = [e[0] for e in m.entries if e[5][0] in ("constant", "open")]
        text = f"{m.name}: {len(board)} of {len(m.entries)} defines traced to FPGA balls"
        if internal:
            text += f"; FPGA-internal: {', '.join(internal)}"
        if idle:
            text += f"; constant or unconnected in the FPGA: {', '.join(idle)}"
        if m.unmapped:
            text += f"; header groups with no CoreGPIO configured: {', '.join(m.unmapped)}"
        yield Finding("FW013", text, refs=[m.spec["fpga"]])
